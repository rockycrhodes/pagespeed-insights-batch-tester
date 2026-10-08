#!/usr/bin/env python3
"""
Generate a formatted Excel results spreadsheet from psi_batch.py output.

Usage:
    python generate_xlsx.py --input psi_results.json --output CWV_Results.xlsx [--label "Pre-launch"]
"""

import argparse
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not found. Install with: pip install openpyxl --break-system-packages")
    sys.exit(1)

# ── Threshold helpers ─────────────────────────────────────────────────────────

def lcp_color(ms):
    if ms is None: return None
    if ms <= 2500: return "C6EFCE"   # green
    if ms <= 4000: return "FFEB9C"   # yellow
    return "FFC7CE"                   # red

def inp_color(ms):
    if ms is None: return None
    if ms <= 200:  return "C6EFCE"
    if ms <= 500:  return "FFEB9C"
    return "FFC7CE"

def cls_color(val):
    if val is None: return None
    if val <= 0.1:  return "C6EFCE"
    if val <= 0.25: return "FFEB9C"
    return "FFC7CE"

def tbt_color(ms):
    if ms is None: return None
    if ms <= 200:  return "C6EFCE"
    if ms <= 600:  return "FFEB9C"
    return "FFC7CE"

def score_color(score):
    if score is None: return None
    if score >= 90: return "C6EFCE"
    if score >= 50: return "FFEB9C"
    return "FFC7CE"

# ── Style helpers ─────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
ALT_FILL = PatternFill("solid", fgColor="EBF3FB")

THIN = Side(style="thin", color="BFBFBF")
ALL_BORDERS = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def header_cell(ws, row, col, value, width=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = ALL_BORDERS
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c

def subheader_cell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    c.fill = SUBHEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = ALL_BORDERS
    return c

def data_cell(ws, row, col, value, fmt=None, color=None, bold=False, url_col=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", size=9, bold=bold)
    c.alignment = Alignment(horizontal="left" if url_col else "center", vertical="center")
    c.border = ALL_BORDERS
    if color:
        c.fill = PatternFill("solid", fgColor=color)
    elif row % 2 == 0:
        c.fill = ALT_FILL
    if fmt:
        c.number_format = fmt
    return c

def fmt_ms(val):
    if val is None: return "N/A"
    if val >= 1000: return f"{val/1000:.2f}s"
    return f"{round(val)}ms"

def fmt_cls(val):
    return "N/A" if val is None else f"{val:.3f}"

def fmt_score(val):
    return "N/A" if val is None else str(int(val))

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input",  required=True, help="JSON output from psi_batch.py")
    parser.add_argument("--output", default="CWV_Results.xlsx")
    parser.add_argument("--label",  default="", help="Optional label e.g. 'Pre-launch' or 'Post-launch'")
    args = parser.parse_args()

    with open(args.input) as f:
        data = json.load(f)

    results = data.get("results", [])
    meta    = data.get("meta", {})

    wb = Workbook()

    # ── Sheet 1: Results ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "CWV Results"
    ws.freeze_panes = "C4"
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 36

    # Title row
    title_text = f"Core Web Vitals Test Results"
    if args.label:
        title_text += f" — {args.label}"
    ts = meta.get("timestamp", "")
    if ts:
        title_text += f"  |  {ts[:10]}"

    ws.merge_cells("A1:O1")
    title = ws["A1"]
    title.value = title_text
    title.font = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E79")
    title.alignment = Alignment(horizontal="left", vertical="center")
    title.border = ALL_BORDERS

    # Meta row
    runs_label = f"Runs per URL: {meta.get('runs', '?')} | Strategies: {', '.join(meta.get('strategies', []))}"
    ws.merge_cells("A2:O2")
    meta_cell = ws["A2"]
    meta_cell.value = runs_label
    meta_cell.font = Font(name="Arial", size=9, color="FFFFFF")
    meta_cell.fill = PatternFill("solid", fgColor="2E75B6")
    meta_cell.alignment = Alignment(horizontal="left", vertical="center")
    meta_cell.border = ALL_BORDERS

    # Group headers row 3
    # Col layout: A=URL, B=Device | C-G=Lab (median) | H-L=Field (CrUX)
    for col, val, w in [
        (1, "URL",    48),
        (2, "Device",  9),
    ]:
        header_cell(ws, 3, col, val, w)

    ws.merge_cells("C3:G3")
    lab_hdr = ws.cell(row=3, column=3, value="LAB DATA (Lighthouse median)")
    lab_hdr.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    lab_hdr.fill = PatternFill("solid", fgColor="2E75B6")
    lab_hdr.alignment = Alignment(horizontal="center", vertical="center")
    lab_hdr.border = ALL_BORDERS

    ws.merge_cells("H3:L3")
    field_hdr = ws.cell(row=3, column=8, value="FIELD DATA (CrUX — real users, 28-day window)")
    field_hdr.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    field_hdr.fill = PatternFill("solid", fgColor="1F5C87")
    field_hdr.alignment = Alignment(horizontal="center", vertical="center")
    field_hdr.border = ALL_BORDERS

    # Sub-headers row 4
    ws.row_dimensions[4].height = 36
    sub_headers = [
        (1, "URL"),
        (2, "Device"),
        (3, "LCP\n(lab)"),
        (4, "TBT\n(lab proxy for INP)"),
        (5, "CLS\n(lab)"),
        (6, "FCP\n(lab)"),
        (7, "Perf\nScore"),
        (8, "LCP\n(field p75)"),
        (9, "INP\n(field p75)"),
        (10, "CLS\n(field p75)"),
        (11, "FCP\n(field p75)"),
        (12, "Overall\nCategory"),
    ]
    for col, val in sub_headers:
        subheader_cell(ws, 4, col, val)

    col_widths = {3: 12, 4: 14, 5: 10, 6: 10, 7: 8, 8: 12, 9: 12, 10: 10, 11: 12, 12: 14}
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # Data rows starting at 5
    for i, r in enumerate(results):
        row = i + 5
        ws.row_dimensions[row].height = 16
        lab = r.get("lab_median", {})
        fld = r.get("field", {})
        url = r.get("url", "")
        strategy = r.get("strategy", "")
        alt = row % 2 == 0

        lcp_lab  = lab.get("lcp_ms")
        tbt_lab  = lab.get("tbt_ms")
        cls_lab  = lab.get("cls")
        fcp_lab  = lab.get("fcp_ms")
        score    = lab.get("perf_score")
        lcp_fld  = fld.get("lcp_p75_ms")
        inp_fld  = fld.get("inp_p75_ms")
        cls_fld  = fld.get("cls_p75")
        fcp_fld  = fld.get("fcp_p75_ms")
        overall  = fld.get("overall", "—")

        # (col, value, color, is_url_col)
        row_data = [
            (1,  url,                  None,                   True),
            (2,  strategy.title(),     None,                   False),
            (3,  fmt_ms(lcp_lab),      lcp_color(lcp_lab),    False),
            (4,  fmt_ms(tbt_lab),      tbt_color(tbt_lab),    False),
            (5,  fmt_cls(cls_lab),     cls_color(cls_lab),    False),
            (6,  fmt_ms(fcp_lab),      None,                   False),
            (7,  fmt_score(score),     score_color(score),    False),
            (8,  fmt_ms(lcp_fld),      lcp_color(lcp_fld),    False),
            (9,  fmt_ms(inp_fld),      inp_color(inp_fld),    False),
            (10, fmt_cls(cls_fld),     cls_color(cls_fld),    False),
            (11, fmt_ms(fcp_fld),      None,                   False),
            (12, overall or "—",       None,                   False),
        ]

        for col, val, color, is_url_col in row_data:
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(horizontal="left" if is_url_col else "center", vertical="center")
            c.border = ALL_BORDERS
            if color:
                c.fill = PatternFill("solid", fgColor=color)
            elif alt and col > 2:
                c.fill = ALT_FILL

    # ── Sheet 2: Legend ───────────────────────────────────────────────────────
    leg = wb.create_sheet("Legend & Thresholds")
    leg.column_dimensions["A"].width = 28
    leg.column_dimensions["B"].width = 20
    leg.column_dimensions["C"].width = 20
    leg.column_dimensions["D"].width = 20

    def leg_header(row, text):
        for col in range(1, 5):
            c = leg.cell(row=row, column=col, value=text if col == 1 else "")
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.border = ALL_BORDERS
        leg.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

    def leg_row(row, metric, good, ni, poor):
        vals = [metric, good, ni, poor]
        colors = [None, "C6EFCE", "FFEB9C", "FFC7CE"]
        for col, (val, color) in enumerate(zip(vals, colors), 1):
            c = leg.cell(row=row, column=col, value=val)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = ALL_BORDERS
            if color:
                c.fill = PatternFill("solid", fgColor=color)

    def leg_subhdr(row):
        for col, val in enumerate(["Metric", "Good ✅", "Needs Improvement ⚠️", "Poor ❌"], 1):
            c = leg.cell(row=row, column=col, value=val)
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="2E75B6")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = ALL_BORDERS

    leg_header(1, "Core Web Vitals Thresholds")
    leg_subhdr(2)
    leg_row(3, "LCP (Largest Contentful Paint)", "≤ 2.5s", "2.5s – 4.0s", "> 4.0s")
    leg_row(4, "INP (Interaction to Next Paint)", "≤ 200ms", "200ms – 500ms", "> 500ms")
    leg_row(5, "CLS (Cumulative Layout Shift)", "≤ 0.1", "0.1 – 0.25", "> 0.25")
    leg_row(6, "TBT (lab proxy for INP)", "≤ 200ms", "200ms – 600ms", "> 600ms")
    leg_row(7, "Performance Score", "≥ 90", "50 – 89", "< 50")

    leg.append([])
    leg_header(9, "Notes")
    notes = [
        (10, "Lab data", "Synthetic Lighthouse test. Use for pre/post comparison. Does not reflect real user experience directly."),
        (11, "Field data (CrUX)", "Aggregated from real Chrome users, past 28 days. May show N/A for low-traffic pages."),
        (12, "TBT vs INP", "TBT is the lab proxy for INP. If field INP is available, prefer that for diagnosis."),
        (13, "Median runs", "Each lab score is the median of multiple runs. Run count is shown in the Results sheet header."),
        (14, "INP replaced FID", "Since March 2024, INP is the official CWV metric. Ignore any FID references in older reports."),
    ]
    for row, label, note_text in notes:
        c = leg.cell(row=row, column=1, value=label)
        c.font = Font(name="Arial", bold=True, size=10)
        c.border = ALL_BORDERS
        n = leg.cell(row=row, column=2, value=note_text)
        n.font = Font(name="Arial", size=10)
        n.alignment = Alignment(wrap_text=True)
        n.border = ALL_BORDERS
        leg.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        leg.row_dimensions[row].height = 28

    wb.save(args.output)
    print(f"Spreadsheet saved: {args.output}")


if __name__ == "__main__":
    main()
